"""Excel store for the hidden Derja labelling desk (/derja).

The workbook has three columns, in this order: text, intent, script. Rows are appended in place, so the
file always matches what the labelling desk shows. Every read and write runs under one lock, and saves go
through a temp file, so a crash never leaves a half-written workbook.
"""
import os
import re
import threading
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .security import clean_text

COLUMNS = ("text", "intent", "script")
SCRIPTS = ("arabic", "arabizi", "mixed")
INTENT_RE = re.compile(r"^[a-z][a-z0-9_]{1,39}$")
ARABIC_RE = re.compile(r"[؀-ۿ]")
LATIN_RE = re.compile(r"[A-Za-z]")

_lock = threading.Lock()


class LabelerError(Exception):
    """Bad input from the desk. The message is safe to show to the person typing."""


class DuplicateError(LabelerError):
    """The same message is already in the file."""


def detect_script(text):
    """Arabic letters only = arabic, Latin letters only = arabizi (Derja written with Latin letters), both = mixed."""
    has_arabic = bool(ARABIC_RE.search(text))
    has_latin = bool(LATIN_RE.search(text))
    if has_arabic and has_latin:
        return "mixed"
    return "arabic" if has_arabic else "arabizi"


def normalize_intent(raw):
    """'Order Status' -> 'order_status'. Returns None when the label cannot be a clean intent name."""
    s = re.sub(r"[\s\-]+", "_", str(raw or "").strip().lower())
    return s if INTENT_RE.match(s) else None


def _dedupe_key(text):
    return " ".join(text.split()).casefold()


def _open(path):
    """Return (workbook, sheet). Creates the file with its header row when it does not exist yet."""
    path = Path(path)
    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "messages"
        ws.append(list(COLUMNS))
        _save(wb, path)
        return wb, ws
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    header = tuple(str(c.value or "").strip().lower() for c in ws[1][:len(COLUMNS)])
    if header != COLUMNS:
        raise LabelerError("The Excel file's first row must be: text, intent, script")
    return wb, ws


def _save(wb, path):
    path = Path(path)
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)


def _rows(ws):
    out = []
    for row in ws.iter_rows(min_row=2, max_col=len(COLUMNS), values_only=True):
        text, intent, script = (str(v or "").strip() for v in row)
        if text:
            out.append({"text": text, "intent": intent, "script": script})
    return out


def read_rows(path):
    with _lock:
        wb, ws = _open(path)
        return _rows(ws)


def export_bytes(path):
    """The current workbook as bytes, for download."""
    with _lock:
        _open(path)
        return Path(path).read_bytes()


def add_row(path, text, intent, script):
    """Validate one labelled message and append it. Returns the stored row."""
    msg = clean_text(text, 1000)
    if len(msg) < 2:
        raise LabelerError("Write the message first.")
    if not (ARABIC_RE.search(msg) or LATIN_RE.search(msg)):
        raise LabelerError("The message needs some letters in it.")

    label = normalize_intent(intent)
    if not label:
        raise LabelerError("Intent must be 2 to 40 letters, digits or underscores, e.g. order_status.")

    script_value = str(script or "").strip().lower()
    if script_value in ("", "auto"):
        script_value = detect_script(msg)
    elif script_value not in SCRIPTS:
        raise LabelerError("Script must be arabic, arabizi or mixed.")

    key = _dedupe_key(msg)
    with _lock:
        wb, ws = _open(path)
        if any(_dedupe_key(r["text"]) == key for r in _rows(ws)):
            raise DuplicateError("This message is already in the file.")
        ws.append([msg, label, script_value])
        # Force text cells to string type so a message starting with "=" is never stored as a formula.
        last = ws.max_row
        for col in range(1, len(COLUMNS) + 1):
            ws.cell(row=last, column=col).data_type = "s"
        _save(wb, path)
    return {"text": msg, "intent": label, "script": script_value}

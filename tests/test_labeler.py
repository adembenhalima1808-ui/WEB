"""Labelling desk (/derja): Excel store unit tests and endpoint tests through Starlette's test client."""
import os
import shutil
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("COOKIE_SECURE", "0")
os.environ.setdefault("SECRET_KEY", "test-secret")
os.environ.setdefault("ENV_FILE", "")

from openpyxl import Workbook, load_workbook
from starlette.testclient import TestClient

from app import labeler, main, security, settings

PASS = "derja-desk-passphrase"


class StoreTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.path = Path(self.dir) / "messages.xlsx"

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_detect_script(self):
        self.assertEqual(labeler.detect_script("وين وصلت"), "arabic")
        self.assertEqual(labeler.detect_script("win wslet commandti?"), "arabizi")
        self.assertEqual(labeler.detect_script("Salut, ma commande وصلت"), "mixed")

    def test_normalize_intent(self):
        self.assertEqual(labeler.normalize_intent(" Order Status "), "order_status")
        self.assertEqual(labeler.normalize_intent("delivery-delay"), "delivery_delay")
        self.assertIsNone(labeler.normalize_intent("1bad"))
        self.assertIsNone(labeler.normalize_intent("x"))
        self.assertIsNone(labeler.normalize_intent("drop;table"))

    def test_add_row_creates_header_and_appends(self):
        row = labeler.add_row(self.path, "وين وصلت طلبيتي؟", "order_status", "auto")
        self.assertEqual(row, {"text": "وين وصلت طلبيتي؟", "intent": "order_status", "script": "arabic"})
        ws = load_workbook(self.path).worksheets[0]
        self.assertEqual([c.value for c in ws[1]], ["text", "intent", "script"])
        self.assertEqual([c.value for c in ws[2]], ["وين وصلت طلبيتي؟", "order_status", "arabic"])

    def test_rejects_bad_input(self):
        with self.assertRaises(labeler.LabelerError):
            labeler.add_row(self.path, " ", "order_status", "auto")
        with self.assertRaises(labeler.LabelerError):
            labeler.add_row(self.path, "123 456", "order_status", "auto")
        with self.assertRaises(labeler.LabelerError):
            labeler.add_row(self.path, "salut", "!!", "auto")
        with self.assertRaises(labeler.LabelerError):
            labeler.add_row(self.path, "salut", "greeting", "klingon")

    def test_duplicates_ignore_case_and_spacing(self):
        labeler.add_row(self.path, "Win commandti?", "order_status", "auto")
        with self.assertRaises(labeler.DuplicateError):
            labeler.add_row(self.path, "  win   COMMANDTI? ", "order_status", "auto")

    def test_formula_looking_text_is_stored_as_text(self):
        labeler.add_row(self.path, "=HYPERLINK(\"http://x\",\"hi\")", "greeting", "arabizi")
        cell = load_workbook(self.path).worksheets[0]["A2"]
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("="))

    def test_refuses_a_file_with_other_columns(self):
        wb = Workbook()
        wb.active.append(["message", "label"])
        wb.save(self.path)
        with self.assertRaises(labeler.LabelerError):
            labeler.read_rows(self.path)

    def test_export_returns_a_workbook(self):
        labeler.add_row(self.path, "salut ça va", "greeting", "arabizi")
        data = labeler.export_bytes(self.path)
        self.assertEqual(data[:2], b"PK")          # xlsx files are zip archives
        self.assertEqual(len(labeler.read_rows(self.path)), 1)


class EndpointTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.saved = (settings.LABELER_PASSPHRASE, settings.DERJA_XLSX)
        settings.LABELER_PASSPHRASE = PASS
        settings.DERJA_XLSX = Path(self.dir) / "messages.xlsx"
        security.limiter.hits.clear()
        self.client = TestClient(main.app)

    def tearDown(self):
        settings.LABELER_PASSPHRASE, settings.DERJA_XLSX = self.saved
        shutil.rmtree(self.dir, ignore_errors=True)

    def login(self, passphrase=PASS):
        return self.client.post("/api/derja/login", json={"passphrase": passphrase})

    def test_page_is_served_and_not_indexed(self):
        r = self.client.get("/derja")
        self.assertEqual(r.status_code, 200)
        self.assertIn('content="noindex, nofollow"', r.text)

    def test_public_home_does_not_link_to_the_desk(self):
        home = Path(main.settings.FRONTEND_DIR, "index.html").read_text(encoding="utf-8")
        self.assertNotIn("derja", home.lower())

    def test_wrong_passphrase_is_refused(self):
        self.assertEqual(self.login("guess").status_code, 401)
        self.assertEqual(self.client.get("/api/derja/state").status_code, 401)

    def test_desk_is_disabled_without_a_passphrase(self):
        settings.LABELER_PASSPHRASE = ""
        self.assertEqual(self.login().status_code, 403)

    def test_add_and_state_after_login(self):
        self.assertEqual(self.login().status_code, 200)
        r = self.client.post("/api/derja/add", json={"text": "win wslet commandti?", "intent": "order_status", "script": "auto"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["row"]["script"], "arabizi")
        state = self.client.get("/api/derja/state").json()
        self.assertEqual(state["count"], 1)
        self.assertEqual(state["intents"], ["order_status"])
        self.assertEqual(state["recent"][0]["text"], "win wslet commandti?")

    def test_duplicate_returns_409_and_bad_input_400(self):
        self.login()
        body = {"text": "salut ça va", "intent": "greeting", "script": "arabizi"}
        self.assertEqual(self.client.post("/api/derja/add", json=body).status_code, 200)
        self.assertEqual(self.client.post("/api/derja/add", json=body).status_code, 409)
        bad = dict(body, intent="!!")
        self.assertEqual(self.client.post("/api/derja/add", json=bad).status_code, 400)

    def test_add_requires_login(self):
        r = self.client.post("/api/derja/add", json={"text": "salut", "intent": "greeting"})
        self.assertEqual(r.status_code, 401)
        self.assertFalse(settings.DERJA_XLSX.exists())

    def test_export_downloads_the_workbook(self):
        self.login()
        self.client.post("/api/derja/add", json={"text": "salut ça va", "intent": "greeting", "script": "arabizi"})
        r = self.client.get("/api/derja/export")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r.headers["content-disposition"])
        self.assertEqual(r.content[:2], b"PK")

    def test_labeller_session_is_not_a_site_role(self):
        self.login()
        me = self.client.get("/api/auth/me").json()
        self.assertIsNone(me["role"])
        self.assertFalse(me["init"])

    def test_logout_ends_the_desk_session(self):
        self.login()
        self.client.post("/api/auth/logout", json={})
        self.assertEqual(self.client.get("/api/derja/state").status_code, 401)


if __name__ == "__main__":
    unittest.main()
